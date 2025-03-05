import os
import requests
from rag import RAG
import docx
from dotenv import load_dotenv

load_dotenv()

RAG_FILE = "rag_data.pkl"
API_URL = "https://api.openai.com/v1/chat/completions"
API_KEY = os.getenv("OPENAI_API_KEY")

def get_chatgpt_response(rol, prompt):
    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json"
    }
    data = {
        "model": "gpt-4o",
        "messages": [
            {"role": "system", "content": rol},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.7,
        "max_tokens": 1024
    }
    response = requests.post(API_URL, headers=headers, json=data)
    if response.status_code == 200:
        return response.json()["choices"][0]["message"]["content"]
    else:
        return f"Error: {response.status_code}, {response.text}"

def save_responses_to_word(responses):
    fileName = "../output/respostes_chatgpt.docx"
    
    if os.path.exists(fileName):
        os.remove(fileName)
    
    doc = docx.Document()
    doc.add_heading("Respostes de ChatGPT", level=1)
    
    for i, response in enumerate(responses, 1):
        doc.add_heading(f"Pregunta {i}", level=2)
        doc.add_paragraph(response)
    
    doc.save(fileName)
    print(f"Respostes desades a {fileName}")

def main():
    if os.path.exists(RAG_FILE):
        rag_system = RAG.load(RAG_FILE)
    else:
        print("RAG file not found!")
        return
    
    responses = []
    
    for i in range(10):
        tema = input(f"Introdueix el tema per a la pregunta {i+1}: ")
        relevant_sentence = rag_system.retrieve(tema)
        
        rol = "Ets un generador d'enunciats d'una assignatura de programació. Només has de generar el text de l'enunciat sense explicacions prèvies ni posteriors. L'enunciat ha de ser clar i concís i no es poden demanar coses que no estiguin explicades en el prompt que se't proporciona. Et donaré diversos exemples per tal que et basis en ells per generar la pregunta. Els exemples contenen la resposta, però només has de generar la pregunta."
        
        meta_prompt = f"El contingut sobre el qual pots fer preguntes és el següent: {relevant_sentence}."
        meta_prompt += "\nGenera les preguntes basant-te en els següents exemples:\n"
        
        for filename in os.listdir("../input/"):
            if filename.endswith(".txt"):
                with open(f"../input/{filename}", "r", encoding="utf-8") as file:
                    meta_prompt += file.read() + "\n"
        
        print(meta_prompt)
        chatgpt_response = get_chatgpt_response(rol, meta_prompt)
        responses.append(chatgpt_response)
    
    save_responses_to_word(responses)

if __name__ == "__main__":
    main()